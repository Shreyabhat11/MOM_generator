"""Generates the small, synthetic fixture files used by the test suite and
the (optional) manual evaluation pass. Run with:
    python fixtures/generate_fixtures.py
Fixtures are intentionally synthetic (no real meeting content) - see spec
section 29 (do not commit example documents with real private information).
"""
from __future__ import annotations

import os

import cv2
import numpy as np
from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

HERE = os.path.dirname(__file__)


def make_typed_notes_txt():
    text = (
        "Weekly Sync - Engineering Team\n"
        "Date: March 14, 2025\n"
        "Location: Conference Room B\n\n"
        "Attendees: Raj, Sarah, Mike\n\n"
        "Discussion: Raj will send revised budget by Friday. "
        "Sarah will review it Monday. Decision: proceed with vendor A.\n\n"
        "Action Items:\n"
        "Raj - prepare financial report - due Friday\n"
        "Sarah - review budget - due Monday\n"
    )
    path = os.path.join(HERE, "typed_notes", "weekly_sync.txt")
    with open(path, "w") as f:
        f.write(text)
    return path


def make_native_pdf():
    path = os.path.join(HERE, "typed_notes", "weekly_sync.pdf")
    c = canvas.Canvas(path, pagesize=LETTER)
    lines = [
        "WEEKLY SYNC",
        "Date: March 14, 2025",
        "Location: Conference Room B",
        "",
        "Attendees: Raj, Sarah, Mike",
        "",
        "- Raj will send revised budget by Friday.",
        "- Sarah will review it Monday.",
        "- Decision: proceed with vendor A.",
    ]
    y = 750
    for line in lines:
        c.drawString(72, y, line)
        y -= 18
    c.save()
    return path


def make_empty_pdf():
    path = os.path.join(HERE, "expected", "empty.pdf")
    c = canvas.Canvas(path, pagesize=LETTER)
    c.save()  # zero content pages... reportlab always emits at least 1 page; used for "near empty" test
    return path


def make_scanned_pdf():
    """A PDF whose only page is a rasterized image (no extractable text
    layer) - exercises the scanned-PDF -> vision-OCR fallback path."""
    import fitz

    img = np.full((600, 800, 3), 255, dtype=np.uint8)
    cv2.putText(img, "SCANNED PAGE (image only)", (40, 300), cv2.FONT_HERSHEY_SIMPLEX, 1, (0, 0, 0), 2)
    img_path = os.path.join(HERE, "scanned_pdfs", "_tmp.png")
    cv2.imwrite(img_path, img)

    doc = fitz.open()
    page = doc.new_page(width=800, height=600)
    page.insert_image(page.rect, filename=img_path)
    out_path = os.path.join(HERE, "scanned_pdfs", "scanned_notes.pdf")
    doc.save(out_path)
    doc.close()
    os.remove(img_path)
    return out_path


def make_docx():
    path = os.path.join(HERE, "typed_notes", "weekly_sync.docx")
    doc = Document()
    doc.add_heading("Weekly Sync", level=1)
    doc.add_paragraph("Date: March 14, 2025")
    doc.add_paragraph("Attendees:")
    for name in ["Raj", "Sarah", "Mike"]:
        doc.add_paragraph(name, style="List Bullet")

    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Action", "Owner", "Deadline"
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = "Prepare report", "Raj", "Friday"

    doc.save(path)
    return path


def make_skewed_image():
    img = np.full((400, 600, 3), 255, dtype=np.uint8)
    cv2.putText(img, "Handwritten-style note", (30, 200), cv2.FONT_HERSHEY_SIMPLEX, 1, (0, 0, 0), 2)
    center = (300, 200)
    matrix = cv2.getRotationMatrix2D(center, 8, 1.0)  # rotate 8 degrees to simulate skew
    rotated = cv2.warpAffine(img, matrix, (600, 400), borderValue=(255, 255, 255))
    path = os.path.join(HERE, "handwritten_notes", "skewed_note.png")
    cv2.imwrite(path, rotated)
    return path


def make_low_contrast_image():
    img = np.full((400, 600, 3), 180, dtype=np.uint8)
    cv2.putText(img, "Low contrast note", (30, 200), cv2.FONT_HERSHEY_SIMPLEX, 1, (160, 160, 160), 2)
    path = os.path.join(HERE, "handwritten_notes", "low_contrast_note.png")
    cv2.imwrite(path, img)
    return path


def make_clean_image():
    img = np.full((400, 600, 3), 255, dtype=np.uint8)
    cv2.putText(img, "Clean typed note", (30, 200), cv2.FONT_HERSHEY_SIMPLEX, 1, (0, 0, 0), 2)
    path = os.path.join(HERE, "typed_notes", "clean_note.png")
    cv2.imwrite(path, img)
    return path


if __name__ == "__main__":
    for f in (
        make_typed_notes_txt,
        make_native_pdf,
        make_empty_pdf,
        make_scanned_pdf,
        make_docx,
        make_skewed_image,
        make_low_contrast_image,
        make_clean_image,
    ):
        print(f())
