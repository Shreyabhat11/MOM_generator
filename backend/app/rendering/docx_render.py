"""Renders a validated MeetingMinutes into a real, professionally formatted
.docx (spec section 18) using python-docx. No plain-text-pretending-to-be-
Word output."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.models.schemas import MeetingMinutes

_ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def _add_heading(doc: Document, text: str) -> None:
    heading = doc.add_heading(text, level=1)
    for run in heading.runs:
        run.font.color.rgb = _ACCENT


def _field_line(doc: Document, label: str, value: str | None) -> None:
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label}: ")
    run_label.bold = True
    p.add_run(value if value else "Not mentioned")


def render_docx(minutes: MeetingMinutes) -> bytes:
    doc = Document()

    title = doc.add_heading(minutes.meeting_title or "Minutes of Meeting", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _field_line(doc, "Date", minutes.date)
    _field_line(doc, "Time", minutes.time)
    _field_line(doc, "Location", minutes.location)
    _field_line(doc, "Organizer", minutes.organizer)

    _add_heading(doc, "Attendees")
    if minutes.attendees:
        for a in minutes.attendees:
            doc.add_paragraph(a, style="List Bullet")
    else:
        doc.add_paragraph("Not mentioned")

    _add_heading(doc, "Agenda")
    if minutes.agenda:
        for a in minutes.agenda:
            doc.add_paragraph(a, style="List Number")
    else:
        doc.add_paragraph("Not mentioned")

    _add_heading(doc, "Discussion Summary")
    narrative = minutes.field_confidence.get("discussion_narrative")
    if narrative:
        doc.add_paragraph(narrative)
    elif minutes.discussion_points:
        for d in minutes.discussion_points:
            doc.add_paragraph(d, style="List Bullet")
    else:
        doc.add_paragraph("Not mentioned")

    _add_heading(doc, "Decisions")
    if minutes.decisions:
        for d in minutes.decisions:
            doc.add_paragraph(d, style="List Bullet")
    else:
        doc.add_paragraph("Not mentioned")

    _add_heading(doc, "Action Items")
    if minutes.action_items:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Action", "Owner", "Deadline"
        for item in minutes.action_items:
            row = table.add_row().cells
            row[0].text = item.task
            row[1].text = item.owner or "Not mentioned"
            row[2].text = item.deadline or "Not mentioned"
    else:
        doc.add_paragraph("Not mentioned")

    _add_heading(doc, "Risks")
    if minutes.risks:
        for r in minutes.risks:
            doc.add_paragraph(r, style="List Bullet")
    else:
        doc.add_paragraph("Not mentioned")

    _add_heading(doc, "Open Questions")
    if minutes.open_questions:
        for q in minutes.open_questions:
            doc.add_paragraph(q, style="List Bullet")
    else:
        doc.add_paragraph("Not mentioned")

    _add_heading(doc, "Next Meeting")
    doc.add_paragraph(minutes.next_meeting or "Not mentioned")

    _add_heading(doc, "Summary")
    doc.add_paragraph(minutes.summary or "Not available")

    core_props = doc.core_properties
    core_props.title = minutes.meeting_title or "Minutes of Meeting"
    core_props.subject = "Minutes of Meeting"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
