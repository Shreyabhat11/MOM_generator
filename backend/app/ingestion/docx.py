"""DOCX ingestion. Unlike the original prototype (which joined all
paragraph text with spaces, destroying structure), this walks the document
body in order and preserves headings, list items, and tables as typed
blocks (spec section 6)."""
from __future__ import annotations

import io
from typing import List

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.models.schemas import BlockType, DocumentKind, ExtractionResult, LayoutBlock, PageResult, TableData


def _classify_paragraph(paragraph: Paragraph) -> BlockType:
    style = (paragraph.style.name or "").lower() if paragraph.style else ""
    if "heading" in style or "title" in style:
        return BlockType.HEADING
    numbering = paragraph._p.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr/"
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr"
    )
    if numbering is not None:
        return BlockType.NUMBERED_ITEM
    if "list bullet" in style:
        return BlockType.BULLET
    return BlockType.PARAGRAPH


def _table_to_block(table: Table, order: int) -> LayoutBlock:
    rows: List[List[str]] = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    flat_text = "\n".join(" | ".join(r) for r in rows)
    return LayoutBlock(
        page=1,
        block_type=BlockType.TABLE,
        text=flat_text,
        table=TableData(rows=rows),
        order=order,
        source="docx",
    )


def extract_docx(file_bytes: bytes) -> ExtractionResult:
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        return ExtractionResult(
            document_kind=DocumentKind.UNKNOWN,
            pages=[],
            warnings=[f"Could not open DOCX file ({type(exc).__name__})."],
        )

    blocks: List[LayoutBlock] = []
    order = 0

    # python-docx doesn't expose body children in a single ordered list by
    # default; walk the underlying XML body so paragraphs and tables come
    # out in the same order they appear in the document.
    body = document.element.body
    para_map = {p._p: p for p in document.paragraphs}
    table_map = {t._tbl: t for t in document.tables}

    for child in body.iterchildren():
        if child in para_map:
            paragraph = para_map[child]
            text = paragraph.text.strip()
            if not text:
                continue
            block_type = _classify_paragraph(paragraph)
            blocks.append(LayoutBlock(page=1, block_type=block_type, text=text, order=order, source="docx"))
            order += 1
        elif child in table_map:
            blocks.append(_table_to_block(table_map[child], order))
            order += 1

    raw_text = "\n".join(b.text for b in blocks)
    warnings: List[str] = []
    if not raw_text.strip():
        warnings.append("No readable text found in this DOCX file.")

    return ExtractionResult(
        document_kind=DocumentKind.DOCX,
        pages=[PageResult(page=1, blocks=blocks, raw_text=raw_text)],
        warnings=warnings,
    )
