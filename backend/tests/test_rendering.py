import io

from docx import Document as DocxReader
from pypdf import PdfReader

from app.models.schemas import ActionItem, MeetingMinutes
from app.rendering.docx_render import render_docx
from app.rendering.pdf_render import render_pdf


def _sample_minutes():
    return MeetingMinutes(
        meeting_title="Weekly Sync",
        date="March 14, 2025",
        location="Conference Room B",
        attendees=["Raj", "Sarah"],
        decisions=["Proceed with vendor A"],
        action_items=[ActionItem(task="Prepare financial report", owner="Raj", deadline="Friday")],
        summary="The team reviewed the budget and agreed to proceed with vendor A.",
    )


def test_docx_is_a_valid_openxml_document():
    content = render_docx(_sample_minutes())
    doc = DocxReader(io.BytesIO(content))  # raises if not a valid docx
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Weekly Sync" in text


def test_docx_contains_expected_sections():
    content = render_docx(_sample_minutes())
    doc = DocxReader(io.BytesIO(content))
    headings = [p.text for p in doc.paragraphs if p.style and "Heading" in (p.style.name or "")]
    for expected in ["Attendees", "Decisions", "Action Items", "Summary"]:
        assert expected in headings


def test_docx_action_item_table_has_correct_data():
    content = render_docx(_sample_minutes())
    doc = DocxReader(io.BytesIO(content))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "Action"
    assert table.rows[1].cells[1].text == "Raj"


def test_docx_missing_fields_render_as_not_mentioned():
    content = render_docx(MeetingMinutes())
    doc = DocxReader(io.BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Not mentioned" in text


def test_pdf_is_valid_and_has_expected_page():
    content = render_pdf(_sample_minutes())
    reader = PdfReader(io.BytesIO(content))
    assert len(reader.pages) >= 1
    text = reader.pages[0].extract_text()
    assert "Weekly Sync" in text
    assert "Prepare financial report" in text


def test_pdf_missing_fields_render_as_not_mentioned():
    content = render_pdf(MeetingMinutes())
    reader = PdfReader(io.BytesIO(content))
    text = "".join(p.extract_text() for p in reader.pages)
    assert "Not mentioned" in text
